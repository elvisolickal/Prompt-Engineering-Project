"""
Ingestion module — load and preprocess writing samples from various file types.
Supports: .txt, .md, .docx, .pdf, .eml
"""
import os
import re
import email
from pathlib import Path
from dataclasses import dataclass
from typing import List, Optional


# ── Bootstrap NLTK data ────────────────────────────────────────────────────────
import nltk
for _res in ["punkt", "punkt_tab"]:
    try:
        nltk.data.find(f"tokenizers/{_res}")
    except LookupError:
        nltk.download(_res, quiet=True)


@dataclass
class Document:
    filename: str
    content: str
    paragraphs: List[str]
    sentences: List[str]
    word_count: int
    source_type: str  # essay | email | article | social | other


class CorpusIngester:
    """Load and preprocess writing samples."""

    MIN_WORDS = 80  # skip documents shorter than this

    # ── Public API ─────────────────────────────────────────────────────────────

    def load_from_text(
        self, text: str, name: str = "pasted_text", source_type: str = "essay"
    ) -> Optional["Document"]:
        """Create a Document from a raw text string (e.g. Streamlit text_area)."""
        cleaned = self._clean(text)
        if len(cleaned.split()) < self.MIN_WORDS:
            return None
        return self._build_doc(cleaned, name, source_type)

    def load_file(self, filepath: str, source_type: str = "essay") -> Optional["Document"]:
        """Load a single file into a Document."""
        raw = self._read_file(filepath)
        if raw is None:
            return None
        cleaned = self._clean(raw)
        if len(cleaned.split()) < self.MIN_WORDS:
            return None
        return self._build_doc(cleaned, Path(filepath).name, source_type)

    def load_uploaded_bytes(
        self, data: bytes, filename: str, source_type: str = "essay"
    ) -> Optional["Document"]:
        """Load from Streamlit UploadedFile bytes."""
        ext = Path(filename).suffix.lower()
        raw = self._parse_bytes(data, ext, filename)
        if raw is None:
            return None
        cleaned = self._clean(raw)
        if len(cleaned.split()) < self.MIN_WORDS:
            return None
        return self._build_doc(cleaned, filename, source_type)

    def get_corpus_text(self, documents: List["Document"]) -> str:
        return "\n\n".join(d.content for d in documents)

    def get_total_words(self, documents: List["Document"]) -> int:
        return sum(d.word_count for d in documents)

    # ── Internal helpers ───────────────────────────────────────────────────────

    def _build_doc(self, cleaned: str, name: str, source_type: str) -> "Document":
        return Document(
            filename=name,
            content=cleaned,
            paragraphs=self._split_paragraphs(cleaned),
            sentences=self._split_sentences(cleaned),
            word_count=len(cleaned.split()),
            source_type=source_type,
        )

    def _clean(self, text: str) -> str:
        text = re.sub(r"https?://\S+", "", text)
        text = re.sub(r"\S+@\S+\.\S+", "", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r" {2,}", " ", text)
        return text.strip()

    def _split_paragraphs(self, text: str) -> List[str]:
        return [p.strip() for p in text.split("\n\n") if len(p.split()) >= 5]

    def _split_sentences(self, text: str) -> List[str]:
        from nltk.tokenize import sent_tokenize
        return [s.strip() for s in sent_tokenize(text) if len(s.split()) >= 3]

    def _read_file(self, filepath: str) -> Optional[str]:
        ext = Path(filepath).suffix.lower()
        try:
            if ext in (".txt", ".md"):
                return Path(filepath).read_text(encoding="utf-8", errors="ignore")
            if ext == ".docx":
                return self._parse_docx_path(filepath)
            if ext == ".pdf":
                return self._parse_pdf_path(filepath)
            if ext == ".eml":
                return self._parse_eml_path(filepath)
        except Exception:
            return None
        return None

    def _parse_bytes(self, data: bytes, ext: str, filename: str) -> Optional[str]:
        import io
        try:
            if ext in (".txt", ".md"):
                return data.decode("utf-8", errors="ignore")
            if ext == ".docx":
                from docx import Document as DocxDoc
                doc = DocxDoc(io.BytesIO(data))
                return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
            if ext == ".pdf":
                import PyPDF2
                reader = PyPDF2.PdfReader(io.BytesIO(data))
                return "\n".join(p.extract_text() or "" for p in reader.pages)
            if ext == ".eml":
                msg = email.message_from_bytes(data)
                parts = []
                if msg.is_multipart():
                    for part in msg.walk():
                        if part.get_content_type() == "text/plain":
                            parts.append(part.get_payload(decode=True).decode("utf-8", errors="ignore"))
                else:
                    parts.append(msg.get_payload(decode=True).decode("utf-8", errors="ignore"))
                return "\n".join(parts)
        except Exception:
            return None
        return None

    def _parse_docx_path(self, filepath: str) -> Optional[str]:
        from docx import Document as DocxDoc
        doc = DocxDoc(filepath)
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    def _parse_pdf_path(self, filepath: str) -> Optional[str]:
        import PyPDF2
        with open(filepath, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            return "\n".join(p.extract_text() or "" for p in reader.pages)

    def _parse_eml_path(self, filepath: str) -> Optional[str]:
        with open(filepath, "rb") as f:
            msg = email.message_from_binary_file(f)
        if msg.is_multipart():
            parts = [
                p.get_payload(decode=True).decode("utf-8", errors="ignore")
                for p in msg.walk()
                if p.get_content_type() == "text/plain"
            ]
            return "\n".join(parts)
        return msg.get_payload(decode=True).decode("utf-8", errors="ignore")
